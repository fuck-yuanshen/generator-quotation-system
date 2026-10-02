import hashlib
import json
import os
import sqlite3
import time
import uuid
from datetime import datetime
from functools import wraps
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlencode, urljoin, urlparse
from urllib.request import urlopen

from flask import (
    Flask,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv("DB_PATH", str(BASE_DIR / "data.db")))
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", str(BASE_DIR / "static" / "uploads")))
GENERATED_DIR = Path(os.getenv("GENERATED_DIR", str(BASE_DIR / "static" / "generated")))
URL_PREFIX = os.getenv("APP_URL_PREFIX", "").strip().rstrip("/")
ALLOWED_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
SELECTION_MODES = {"single", "multi"}

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("FLASK_SECRET_KEY", "replace-this-with-random-secret")
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

ADMIN_USER = os.getenv("ADMIN_USERNAME", "admin")
ADMIN_PASS = os.getenv("ADMIN_PASSWORD")
PUBLIC_BASE_URL = os.getenv("PUBLIC_BASE_URL", "").strip().rstrip("/")
WECHAT_APP_ID = os.getenv("WECHAT_APP_ID", "").strip()
WECHAT_APP_SECRET = os.getenv("WECHAT_APP_SECRET", "").strip()
WECHAT_SHARE_TITLE = os.getenv("WECHAT_SHARE_TITLE", "发电机组报价系统").strip()
WECHAT_SHARE_DESC = os.getenv("WECHAT_SHARE_DESC", "在线选择配置，快速生成报价单。").strip()
WECHAT_SHARE_IMG_URL = os.getenv("WECHAT_SHARE_IMG_URL", "").strip()
WECHAT_SHARE_TITLE = os.getenv("WECHAT_SHARE_TITLE", "").strip()
WECHAT_SHARE_DESC = os.getenv("WECHAT_SHARE_DESC", "").strip()
DEFAULT_SHARE_TITLE = "\u53d1\u7535\u673a\u7ec4\u8ba1\u7b97\u5668"
DEFAULT_SHARE_DESC = "\u5143\u52a8\u7535\u529b\uff0c\u7535\u4eae\u4e16\u754c\U0001F30D"

WECHAT_CACHE: Dict[str, Dict[str, Any]] = {
    "access_token": {"value": "", "expires_at": 0},
    "jsapi_ticket": {"value": "", "expires_at": 0},
}


class PrefixMiddleware:
    def __init__(self, wsgi_app: Any, prefix: str) -> None:
        self.wsgi_app = wsgi_app
        self.prefix = prefix or ""

    def __call__(self, environ: Dict[str, Any], start_response: Any) -> Any:
        if not self.prefix:
            return self.wsgi_app(environ, start_response)

        path_info = environ.get("PATH_INFO", "") or ""
        if path_info.startswith(self.prefix):
            environ["SCRIPT_NAME"] = self.prefix
            environ["PATH_INFO"] = path_info[len(self.prefix) :] or "/"
            return self.wsgi_app(environ, start_response)

        if path_info in {"", "/"}:
            return self.wsgi_app(environ, start_response)

        start_response("404 Not Found", [("Content-Type", "text/plain; charset=utf-8")])
        return [b"Not Found"]


if URL_PREFIX:
    app.wsgi_app = PrefixMiddleware(app.wsgi_app, URL_PREFIX)


def clean_text_setting(value: str, default: str) -> str:
    text = str(value or "").strip()
    if not text:
        return default
    if set(text) == {"?"}:
        return default
    return text


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        g.db = conn
    return g.db


@app.teardown_appcontext
def close_db(_e: Exception) -> None:
    db = g.pop("db", None)
    if db:
        db.close()


def to_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def to_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def current_timestamp() -> int:
    return int(time.time())


def normalize_mode(value: Any) -> str:
    mode = str(value or "").strip().lower()
    return mode if mode in SELECTION_MODES else "single"


def ensure_column(table_name: str, column_name: str, definition: str) -> None:
    db = get_db()
    cols = [row["name"] for row in db.execute(f"PRAGMA table_info({table_name})").fetchall()]
    if column_name not in cols:
        db.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}")
        db.commit()


def absolute_url(path_or_url: str) -> str:
    value = str(path_or_url or "").strip()
    if not value:
        return ""
    if value.startswith(("http://", "https://")):
        return value

    base = PUBLIC_BASE_URL or request.url_root
    return urljoin(base.rstrip("/") + "/", value.lstrip("/"))


def build_default_share_image_url() -> str:
    if WECHAT_SHARE_IMG_URL:
        return absolute_url(WECHAT_SHARE_IMG_URL)

    generated_candidates = [
        "contract_Q20260313003.png",
        "contract_Q20260313001.png",
        "quotation_Q20260313002.png",
    ]
    for filename in generated_candidates:
        candidate = GENERATED_DIR / filename
        if candidate.exists():
            return absolute_url(f"/generated/{filename}")

    return ""


def build_share_meta() -> Dict[str, str]:
    current_path = request.full_path[:-1] if request.full_path.endswith("?") else request.full_path
    current_url = absolute_url(current_path or request.path or "/")
    image_url = build_default_share_image_url()
    title = clean_text_setting(WECHAT_SHARE_TITLE, DEFAULT_SHARE_TITLE)
    desc = clean_text_setting(WECHAT_SHARE_DESC, DEFAULT_SHARE_DESC)
    return {"title": title, "desc": desc, "link": current_url, "img_url": image_url}
    return {
        "title": WECHAT_SHARE_TITLE or "发电机组报价系统",
        "desc": WECHAT_SHARE_DESC or "在线选择配置，快速生成报价单。",
        "link": current_url,
        "img_url": image_url,
    }


def fetch_json(url: str) -> Dict[str, Any]:
    with urlopen(url, timeout=10) as resp:
        return json.loads(resp.read().decode("utf-8"))


def get_cached_wechat_value(key: str) -> str:
    entry = WECHAT_CACHE.get(key, {})
    value = str(entry.get("value", ""))
    expires_at = to_int(entry.get("expires_at"), 0)
    if value and expires_at > current_timestamp() + 60:
        return value
    return ""


def set_cached_wechat_value(key: str, value: str, expires_in: int) -> str:
    ttl = max(300, to_int(expires_in, 7200) - 120)
    WECHAT_CACHE[key] = {"value": value, "expires_at": current_timestamp() + ttl}
    return value


def get_wechat_access_token() -> str:
    cached = get_cached_wechat_value("access_token")
    if cached:
        return cached
    if not WECHAT_APP_ID or not WECHAT_APP_SECRET:
        raise RuntimeError("未配置微信公众号 AppID 或 AppSecret")

    params = urlencode(
        {
            "grant_type": "client_credential",
            "appid": WECHAT_APP_ID,
            "secret": WECHAT_APP_SECRET,
        }
    )
    data = fetch_json(f"https://api.weixin.qq.com/cgi-bin/token?{params}")
    token = str(data.get("access_token", "")).strip()
    if not token:
        raise RuntimeError(data.get("errmsg") or "获取微信 access_token 失败")
    return set_cached_wechat_value("access_token", token, to_int(data.get("expires_in"), 7200))


def get_wechat_jsapi_ticket() -> str:
    cached = get_cached_wechat_value("jsapi_ticket")
    if cached:
        return cached

    params = urlencode({"access_token": get_wechat_access_token(), "type": "jsapi"})
    data = fetch_json(f"https://api.weixin.qq.com/cgi-bin/ticket/getticket?{params}")
    ticket = str(data.get("ticket", "")).strip()
    if str(data.get("errcode", 0)) != "0" or not ticket:
        raise RuntimeError(data.get("errmsg") or "获取微信 jsapi_ticket 失败")
    return set_cached_wechat_value("jsapi_ticket", ticket, to_int(data.get("expires_in"), 7200))


def build_wechat_signature(target_url: str) -> Dict[str, Any]:
    parsed = urlparse(target_url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise RuntimeError("分享地址不是有效的 HTTP/HTTPS 地址")

    nonce_str = uuid.uuid4().hex[:16]
    timestamp = current_timestamp()
    jsapi_ticket = get_wechat_jsapi_ticket()
    sign_string = (
        f"jsapi_ticket={jsapi_ticket}&noncestr={nonce_str}"
        f"&timestamp={timestamp}&url={target_url}"
    )
    signature = hashlib.sha1(sign_string.encode("utf-8")).hexdigest()
    return {
        "appId": WECHAT_APP_ID,
        "timestamp": timestamp,
        "nonceStr": nonce_str,
        "signature": signature,
        "jsApiList": [
            "updateAppMessageShareData",
            "updateTimelineShareData",
            "onMenuShareAppMessage",
            "onMenuShareTimeline",
        ],
    }


def init_db() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)

    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS categories(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            description TEXT DEFAULT '',
            required INTEGER NOT NULL DEFAULT 1,
            selection_mode TEXT NOT NULL DEFAULT 'single',
            sort_order INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS options(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            description TEXT DEFAULT '',
            price REAL NOT NULL DEFAULT 0,
            image_path TEXT DEFAULT '',
            sort_order INTEGER NOT NULL DEFAULT 0,
            is_active INTEGER NOT NULL DEFAULT 1,
            FOREIGN KEY(category_id) REFERENCES categories(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS quotes(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            quote_no TEXT NOT NULL UNIQUE,
            customer_name TEXT DEFAULT '',
            customer_phone TEXT DEFAULT '',
            company_name TEXT DEFAULT '',
            project_name TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            subtotal REAL NOT NULL DEFAULT 0,
            tax_rate REAL NOT NULL DEFAULT 0.13,
            tax_amount REAL NOT NULL DEFAULT 0,
            total REAL NOT NULL DEFAULT 0,
            selections_json TEXT NOT NULL DEFAULT '[]',
            contract_image_path TEXT DEFAULT '',
            quotation_file_path TEXT DEFAULT '',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS settings(
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL DEFAULT ''
        );
        """
    )
    db.commit()

    # migration for old databases
    ensure_column("categories", "selection_mode", "TEXT NOT NULL DEFAULT 'single'")
    ensure_column("options", "params_text", "TEXT NOT NULL DEFAULT ''")
    ensure_column("quotes", "quotation_image_path", "TEXT DEFAULT ''")
    ensure_column("quotes", "quotation_meta_json", "TEXT DEFAULT ''")

    row = db.execute("SELECT COUNT(*) c FROM categories").fetchone()
    if row["c"] == 0:
        seed_defaults()
    else:
        ensure_optional_multi_category()


def seed_defaults() -> None:
    db = get_db()
    categories = [
        ("机组功率", "按负载选择机组功率", 1, "single", 10),
        ("发动机品牌", "选择发动机品牌", 1, "single", 20),
        ("发电机品牌", "选择发电机品牌", 1, "single", 30),
        ("控制系统", "选择控制柜类型", 1, "single", 40),
        ("可选配置", "可多选，按项目需求增加", 0, "multi", 50),
    ]
    db.executemany(
        "INSERT INTO categories(name,description,required,selection_mode,sort_order) VALUES(?,?,?,?,?)",
        categories,
    )
    rows = db.execute("SELECT id,name FROM categories").fetchall()
    cmap = {r["name"]: r["id"] for r in rows}

    options = {
        "机组功率": [
            ("100kW", 48000, "额定功率:100kW\n额定电压:400V\n频率:50Hz"),
            ("200kW", 86000, "额定功率:200kW\n额定电压:400V\n频率:50Hz"),
            ("300kW", 128000, "额定功率:300kW\n额定电压:400V\n频率:50Hz"),
            ("500kW", 195000, "额定功率:500kW\n额定电压:400V\n频率:50Hz"),
        ],
        "发动机品牌": [
            ("国产", 0, "缸数:6缸\n冷却方式:水冷"),
            ("康明斯", 12000, "缸数:6缸\n冷却方式:水冷\n控制系统:ECU"),
            ("珀金斯", 26000, "缸数:6缸\n冷却方式:水冷\n控制系统:ECU"),
        ],
        "发电机品牌": [
            ("国产", 0, "绝缘等级:H级\n防护等级:IP23"),
            ("马拉松", 9000, "绝缘等级:H级\n防护等级:IP23"),
            ("斯坦福", 15000, "绝缘等级:H级\n防护等级:IP23"),
        ],
        "控制系统": [
            ("基础控制", 0, "显示:LCD\n保护:过压过流"),
            ("ATS自动切换", 6000, "ATS:支持\n切换时间:<10s"),
            ("远程监控", 12000, "远程监控:支持\n通信:4G/网口"),
        ],
        "可选配置": [
            ("静音箱", 18000, "降噪值:15-20dB\n材质:镀锌板"),
            ("并机柜", 8000, "并机台数:2-4台\n控制:自动均载"),
            ("底座油箱", 5000, "容量:8小时油量\n材质:碳钢"),
            ("运输安装", 8000, "服务:运输+吊装+调试"),
        ],
    }

    for category_name, items in options.items():
        for index, (name, price, params_text) in enumerate(items, start=1):
            db.execute(
                """
                INSERT INTO options(category_id,name,price,params_text,sort_order,is_active)
                VALUES(?,?,?,?,?,1)
                """,
                (cmap[category_name], name, price, params_text, index * 10),
            )

    default_settings = {
        "company_name": "某某机电设备有限公司",
        "company_contact": "销售部",
        "company_phone": "138-0000-0000",
        "quotation_show_company": "1",
        "quotation_show_contact": "1",
        "company_address": "",
        "tax_rate": "0.13",
        "contract_notice": "本合同图片为系统生成草稿，正式合同以盖章版为准。",
        "delivery_term": "常规机组交期 7-15 个工作日，定制机组另行确认。",
    }
    for key, value in default_settings.items():
        db.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, value))
    db.commit()


def ensure_optional_multi_category() -> None:
    db = get_db()
    count_row = db.execute("SELECT COUNT(*) c FROM categories WHERE selection_mode='multi'").fetchone()
    if count_row["c"] > 0:
        return

    category_row = db.execute("SELECT id FROM categories WHERE name='可选配置'").fetchone()
    if category_row:
        category_id = category_row["id"]
        db.execute("UPDATE categories SET selection_mode='multi', required=0 WHERE id=?", (category_id,))
    else:
        max_sort_row = db.execute("SELECT COALESCE(MAX(sort_order), 0) AS v FROM categories").fetchone()
        sort_order = to_int(max_sort_row["v"], 0) + 10
        db.execute(
            """
            INSERT INTO categories(name,description,required,selection_mode,sort_order)
            VALUES(?,?,?,?,?)
            """,
            ("可选配置", "可多选，按项目需求增加", 0, "multi", sort_order),
        )
        category_id = db.execute("SELECT id FROM categories WHERE name='可选配置'").fetchone()["id"]

    option_count = db.execute("SELECT COUNT(*) c FROM options WHERE category_id=?", (category_id,)).fetchone()["c"]
    if option_count == 0:
        defaults = [
            ("静音箱", 18000, "降噪值:15-20dB\n材质:镀锌板"),
            ("并机柜", 8000, "并机台数:2-4台\n控制:自动均载"),
            ("底座油箱", 5000, "容量:8小时油量\n材质:碳钢"),
            ("运输安装", 8000, "服务:运输+吊装+调试"),
        ]
        for idx, (name, price, params_text) in enumerate(defaults, start=1):
            db.execute(
                """
                INSERT INTO options(category_id,name,price,params_text,sort_order,is_active)
                VALUES(?,?,?,?,?,1)
                """,
                (category_id, name, price, params_text, idx * 10),
            )
    db.commit()


def get_settings() -> Dict[str, str]:
    rows = get_db().execute("SELECT key,value FROM settings").fetchall()
    data = {r["key"]: r["value"] for r in rows}
    data.setdefault("company_name", "某某机电设备有限公司")
    data.setdefault("company_contact", "销售部")
    data.setdefault("company_phone", "")
    data.setdefault("quotation_show_company", "1")
    data.setdefault("quotation_show_contact", "1")
    data.setdefault("company_address", "")
    data.setdefault("tax_rate", "0.13")
    data.setdefault("contract_notice", "")
    data.setdefault("delivery_term", "常规机组交期 7-15 个工作日，定制机组另行确认。")
    return data


def set_settings(data: Dict[str, str]) -> None:
    db = get_db()
    for key, value in data.items():
        db.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, str(value).strip()))
    db.commit()


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("is_admin"):
            return redirect(url_for("admin_login"))
        return fn(*args, **kwargs)

    return wrapper


def generate_quote_no() -> str:
    day = datetime.now().strftime("Q%Y%m%d")
    row = get_db().execute("SELECT COUNT(*) c FROM quotes WHERE quote_no LIKE ?", (f"{day}%",)).fetchone()
    return f"{day}{row['c'] + 1:03d}"


def parse_selections(raw: Any) -> Dict[int, Any]:
    if not isinstance(raw, dict):
        return {}
    parsed: Dict[int, Any] = {}
    for k, v in raw.items():
        category_id = to_int(k, 0)
        if category_id <= 0:
            continue
        if isinstance(v, list):
            ids = []
            for item in v:
                option_id = to_int(item, 0)
                if option_id > 0 and option_id not in ids:
                    ids.append(option_id)
            parsed[category_id] = ids
        else:
            option_id = to_int(v, 0)
            parsed[category_id] = option_id if option_id > 0 else 0
    return parsed


def parse_params_text(params_text: str) -> List[Dict[str, str]]:
    lines = str(params_text or "").replace("\r\n", "\n").split("\n")
    params = []
    for line in lines:
        text = line.strip()
        if not text:
            continue
        if ":" in text:
            key, value = text.split(":", 1)
        elif "：" in text:
            key, value = text.split("：", 1)
        else:
            key, value = "参数", text
        params.append({"key": key.strip() or "参数", "value": value.strip()})
    return params


def categories_with_options(active_only: bool = True) -> List[Dict[str, Any]]:
    db = get_db()
    categories = db.execute("SELECT * FROM categories ORDER BY sort_order,id").fetchall()
    result = []
    for cat in categories:
        sql = "SELECT * FROM options WHERE category_id=?"
        if active_only:
            sql += " AND is_active=1"
        sql += " ORDER BY sort_order,id"
        options = db.execute(sql, (cat["id"],)).fetchall()
        result.append(
            {
                "id": cat["id"],
                "name": cat["name"],
                "description": cat["description"] or "",
                "required": bool(cat["required"]),
                "selection_mode": normalize_mode(cat["selection_mode"]),
                "sort_order": cat["sort_order"],
                "options": [
                    {
                        "id": op["id"],
                        "category_id": op["category_id"],
                        "name": op["name"],
                        "description": op["description"] or "",
                        "price": float(op["price"]),
                        "image_path": op["image_path"] or "",
                        "params_text": op["params_text"] or "",
                        "params": parse_params_text(op["params_text"] or ""),
                        "sort_order": op["sort_order"],
                        "is_active": bool(op["is_active"]),
                    }
                    for op in options
                ],
            }
        )
    return result


def build_quote_items(selection_map: Dict[int, Any], strict_required: bool = True) -> Tuple[List[Dict[str, Any]], List[str]]:
    db = get_db()
    categories = db.execute(
        "SELECT id,name,required,selection_mode FROM categories ORDER BY sort_order,id"
    ).fetchall()
    items: List[Dict[str, Any]] = []
    errors: List[str] = []

    for cat in categories:
        mode = normalize_mode(cat["selection_mode"])
        raw_value = selection_map.get(cat["id"])

        selected_ids: List[int] = []
        if mode == "multi":
            if isinstance(raw_value, list):
                selected_ids = [to_int(v, 0) for v in raw_value if to_int(v, 0) > 0]
            elif to_int(raw_value, 0) > 0:
                selected_ids = [to_int(raw_value, 0)]
        else:
            selected_id = to_int(raw_value, 0)
            if selected_id > 0:
                selected_ids = [selected_id]

        if strict_required and cat["required"] and not selected_ids:
            errors.append(f"分类「{cat['name']}」为必选项")
            continue
        if not selected_ids:
            continue

        # single mode only keeps first selected value
        if mode == "single" and len(selected_ids) > 1:
            selected_ids = selected_ids[:1]

        for option_id in selected_ids:
            row = db.execute(
                """
                SELECT o.id,o.name,o.price,o.image_path,c.name category_name
                      ,o.params_text
                FROM options o
                JOIN categories c ON c.id=o.category_id
                WHERE o.id=? AND o.category_id=? AND o.is_active=1
                """,
                (option_id, cat["id"]),
            ).fetchone()
            if not row:
                errors.append(f"分类「{cat['name']}」的选项无效")
                continue
            items.append(
                {
                    "category_id": cat["id"],
                    "category_name": row["category_name"],
                    "option_id": row["id"],
                    "option_name": row["name"],
                    "price": float(row["price"]),
                    "image_path": row["image_path"] or "",
                    "params_text": row["params_text"] or "",
                    "params": parse_params_text(row["params_text"] or ""),
                }
            )

    return items, errors


def row_quote(quote_id: int) -> Optional[Dict[str, Any]]:
    row = get_db().execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
    if not row:
        return None
    quote = dict(row)
    quote["subtotal"] = float(quote["subtotal"])
    quote["tax_rate"] = float(quote["tax_rate"])
    quote["tax_amount"] = float(quote["tax_amount"])
    quote["total"] = float(quote["total"])
    quote["selections"] = json.loads(quote["selections_json"] or "[]")
    try:
        quote["quotation_meta"] = json.loads(quote.get("quotation_meta_json") or "{}")
    except (TypeError, ValueError, json.JSONDecodeError):
        quote["quotation_meta"] = {}
    return quote


def resolve_generated_path(rel_path: str) -> Optional[Path]:
    rel = str(rel_path or "").strip()
    if not rel:
        return None
    name = Path(rel).name
    if not name:
        return None
    path = GENERATED_DIR / name
    return path if path.exists() else None


def load_font(size: int, bold: bool = False):
    # 优先使用 Linux 容器中的 CJK 字体，避免中文乱码。
    candidates = (
        [
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
            "C:/Windows/Fonts/msyhbd.ttc",
            "C:/Windows/Fonts/arialbd.ttf",
        ]
        if bold
        else [
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
            "C:/Windows/Fonts/msyh.ttc",
            "C:/Windows/Fonts/arial.ttf",
        ]
    )
    for path in candidates:
        font_path = Path(path)
        if font_path.exists():
            try:
                return ImageFont.truetype(str(font_path), size=size)
            except OSError:
                continue
    return ImageFont.load_default()


def get_selected_item(quote: Dict[str, Any], keywords: List[str]) -> Dict[str, Any]:
    for item in quote.get("selections", []):
        name = str(item.get("category_name", ""))
        if any(k in name for k in keywords):
            return item
    return {}


def get_selected_items(quote: Dict[str, Any], keywords: List[str]) -> List[Dict[str, Any]]:
    result = []
    for item in quote.get("selections", []):
        name = str(item.get("category_name", ""))
        if any(k in name for k in keywords):
            result.append(item)
    return result


def item_params_map(item: Dict[str, Any]) -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    for p in item.get("params", []) or []:
        key = str(p.get("key", "")).strip()
        value = str(p.get("value", "")).strip()
        if key and value:
            mapping[key] = value
    return mapping


def pick_param_value(param_map: Dict[str, str], keys: List[str], default: str = "") -> str:
    for key in keys:
        for actual_key, value in param_map.items():
            if key in actual_key:
                return value
    return default


def number_to_rmb_upper(value: float) -> str:
    nums = "零壹贰叁肆伍陆柒捌玖"
    units = ["", "拾", "佰", "仟"]
    sections = ["", "万", "亿", "万亿"]

    def section_to_text(section_num: int) -> str:
        result = ""
        unit_index = 0
        zero = True
        while section_num > 0:
            num = section_num % 10
            if num == 0:
                if not zero:
                    zero = True
                    result = nums[0] + result
            else:
                zero = False
                result = nums[num] + units[unit_index] + result
            unit_index += 1
            section_num //= 10
        return result

    integer = int(round(value))
    if integer <= 0:
        return "零元整"

    text = ""
    sec_index = 0
    need_zero = False
    while integer > 0:
        section_num = integer % 10000
        if section_num == 0:
            if text and not text.startswith("零"):
                text = "零" + text
        else:
            section_text = section_to_text(section_num) + sections[sec_index]
            if need_zero and not section_text.startswith("零"):
                text = "零" + text
            text = section_text + text
        need_zero = section_num < 1000 and section_num > 0
        integer //= 10000
        sec_index += 1
    text = text.replace("零零", "零").rstrip("零")
    return text + "元整"


def wrap_text_by_width(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> List[str]:
    lines: List[str] = []
    current = ""
    for ch in str(text or ""):
        if ch == "\n":
            lines.append(current)
            current = ""
            continue
        candidate = current + ch
        if draw.textlength(candidate, font=font) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = ch
    if current or not lines:
        lines.append(current)
    return lines


def build_quotation_form_defaults(quote: Dict[str, Any]) -> Dict[str, Any]:
    power_item = get_selected_item(quote, ["机组功率", "功率"])
    engine_item = get_selected_item(quote, ["发动机"])
    alternator_item = get_selected_item(quote, ["发电机"])
    control_item = get_selected_item(quote, ["控制"])
    base_item = get_selected_item(quote, ["底座"])
    color_item = get_selected_item(quote, ["颜色"])
    tank_item = get_selected_item(quote, ["水箱"])
    optional_items = get_selected_items(quote, ["可选", "增选"])

    engine_map = item_params_map(engine_item)
    alternator_map = item_params_map(alternator_item)
    power_map = item_params_map(power_item)

    optional_names = "，".join(item.get("option_name", "") for item in optional_items if item.get("option_name"))

    defaults = {
        "product_series": engine_item.get("option_name", "") or "机组系列",
        "unit_power": power_item.get("option_name", "") or pick_param_value(power_map, ["额定功率", "功率"], ""),
        "unit_name": "台",
        "quantity": "1",
        "unit_price": f"{quote.get('total', 0):.2f}",
        "engine_brand": engine_item.get("option_name", ""),
        "engine_model": pick_param_value(engine_map, ["型号"], engine_item.get("option_name", "")),
        "engine_power": pick_param_value(engine_map, ["功率", "额定功率"], pick_param_value(power_map, ["额定功率", "功率"], "")),
        "engine_emission": pick_param_value(engine_map, ["排放"], ""),
        "tank_type": tank_item.get("option_name", "") or pick_param_value(engine_map, ["水箱"], ""),
        "alternator_type": alternator_item.get("option_name", ""),
        "alternator_model": pick_param_value(alternator_map, ["型号"], alternator_item.get("option_name", "")),
        "alternator_power": pick_param_value(alternator_map, ["功率", "额定功率"], ""),
        "alternator_shell": pick_param_value(alternator_map, ["机壳"], ""),
        "alternator_voltage": pick_param_value(alternator_map, ["电压"], pick_param_value(power_map, ["额定电压"], "")),
        "alternator_frequency": pick_param_value(alternator_map, ["频率"], pick_param_value(power_map, ["频率"], "")),
        "control_system": control_item.get("option_name", ""),
        "base_type": base_item.get("option_name", ""),
        "machine_color": color_item.get("option_name", ""),
        "optional_config": optional_names,
        "other_config": "",
        "payment_term": "预付30%定金，尾款等货到甲方所在地提货市场流站，验货满意后付清。",
        "warranty_term": "质保一年或者累计运行1000小时（两者以先到为准）。",
        "delivery_term": get_settings().get("delivery_term", "交货期：收到定金后5天内发货。"),
        "validity_term": "此报价从即日起1个月内有效。",
    }
    return defaults


def merge_quotation_form(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    result: Dict[str, Any] = {}
    override = override or {}
    for key, default_val in base.items():
        value = override.get(key, default_val)
        result[key] = str(value or "").strip()

    qty = to_int(result.get("quantity"), 1)
    if qty <= 0:
        qty = 1
    result["quantity"] = str(qty)

    price = to_float(result.get("unit_price"), to_float(base.get("unit_price"), 0))
    result["unit_price"] = f"{price:.2f}"
    return result


def generate_contract_image(quote: Dict[str, Any]) -> str:
    out_file = GENERATED_DIR / f"contract_{quote['quote_no']}.png"
    rel = f"/generated/{out_file.name}"

    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    font_title = load_font(52, True)
    font_header = load_font(28, True)
    font_body = load_font(22)

    draw.rectangle([(30, 30), (1210, 140)], fill=(241, 247, 255), outline=(60, 90, 140), width=3)
    title = "设备采购合同（草稿）"
    width = draw.textlength(title, font=font_title)
    draw.text(((1240 - width) / 2, 55), title, fill=(30, 55, 100), font=font_title)

    y = 190
    base_lines = [
        f"合同编号：{quote['quote_no']}",
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"客户：{quote.get('customer_name') or '-'}",
        f"公司：{quote.get('company_name') or '-'}",
        f"电话：{quote.get('customer_phone') or '-'}",
        f"项目：{quote.get('project_name') or '-'}",
    ]
    for line in base_lines:
        draw.text((60, y), line, fill=(45, 45, 45), font=font_body)
        y += 40

    y += 12
    draw.text((60, y), "配置明细", fill=(30, 75, 145), font=font_header)
    y += 50
    for index, item in enumerate(quote["selections"], start=1):
        draw.text(
            (70, y),
            f"{index}. {item['category_name']}：{item['option_name']}  {item['price']:,.2f}",
            fill=(50, 50, 50),
            font=font_body,
        )
        y += 34

    y += 20
    draw.text((60, y), f"未税小计：{quote['subtotal']:,.2f}", fill=(35, 35, 35), font=font_body)
    y += 34
    draw.text((60, y), f"税额（{quote['tax_rate']*100:.0f}%）：{quote['tax_amount']:,.2f}", fill=(35, 35, 35), font=font_body)
    y += 34
    draw.text((60, y), f"含税总价：{quote['total']:,.2f}", fill=(205, 70, 45), font=font_header)

    image.save(out_file, "PNG")
    return rel


def generate_quotation_image(quote: Dict[str, Any], form: Optional[Dict[str, Any]] = None) -> str:
    out_file = GENERATED_DIR / f"quotation_{quote['quote_no']}.png"
    rel = f"/generated/{out_file.name}"
    settings = get_settings()

    defaults = build_quotation_form_defaults(quote)
    form_data = merge_quotation_form(defaults, form or {})
    quantity = to_int(form_data.get("quantity"), 1)
    unit_price = to_float(form_data.get("unit_price"), quote.get("total", 0))
    amount = round(quantity * unit_price, 2)
    amount_upper = number_to_rmb_upper(amount)

    config_lines = [
        (
            f"发动机：{form_data['engine_brand']}（{form_data['product_series']}）系列，"
            f"型号：{form_data['engine_model']}，功率：{form_data['engine_power']}，排放：{form_data['engine_emission']}"
        ).strip("，"),
        f"水箱：{form_data['tank_type']}".strip("："),
        (
            f"发电机：{form_data['alternator_type']}，型号：{form_data['alternator_model']}，"
            f"功率：{form_data['alternator_power']}，机壳：{form_data['alternator_shell']}，"
            f"电压：{form_data['alternator_voltage']}，频率：{form_data['alternator_frequency']}"
        ).strip("，"),
        f"控制系统：{form_data['control_system']}",
        f"底座类型：{form_data['base_type']}",
        f"整机颜色：{form_data['machine_color']}",
    ]
    if form_data.get("optional_config"):
        config_lines.append(f"增配：{form_data['optional_config']}")
    if form_data.get("other_config"):
        config_lines.append(f"其他：{form_data['other_config']}")
    config_lines = [line for line in config_lines if line and line.strip("：")]

    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    font_title = load_font(58, True)
    font_sub = load_font(26, True)
    font_head = load_font(24, True)
    font_body = load_font(22)
    font_note = load_font(20)

    # 标题
    title = "发电机组报价单"
    title_w = draw.textlength(title, font=font_title)
    draw.text(((1240 - title_w) / 2, 80), title, fill=(0, 0, 0), font=font_title)

    # 第一段标题
    y = 190
    draw.text((100, y), "1.配置及价格（单位：人民币元）", fill=(0, 0, 0), font=font_sub)
    y += 48

    # 表格基本参数
    x0 = 100
    table_w = 1040
    col_widths = [150, 160, 120, 120, 170, 170]
    row_h = 46
    table_head_h = 46

    # 表头
    draw.rectangle([(x0, y), (x0 + table_w, y + table_head_h)], outline=(0, 0, 0), width=2)
    headers = ["名称", "功率", "单位", "数量", "单价", "金额"]
    x = x0
    for i, head in enumerate(headers):
        if i > 0:
            draw.line([(x, y), (x, y + table_head_h)], fill=(0, 0, 0), width=2)
        tw = draw.textlength(head, font=font_head)
        draw.text((x + (col_widths[i] - tw) / 2, y + 9), head, fill=(0, 0, 0), font=font_head)
        x += col_widths[i]

    # 主数据行
    y += table_head_h
    draw.rectangle([(x0, y), (x0 + table_w, y + row_h)], outline=(0, 0, 0), width=2)
    row_values = [
        "发电机组",
        form_data.get("unit_power", ""),
        form_data.get("unit_name", "台"),
        str(quantity),
        f"{unit_price:.0f}" if unit_price.is_integer() else f"{unit_price:.2f}",
        f"{amount:.0f}" if amount.is_integer() else f"{amount:.2f}",
    ]
    x = x0
    for i, val in enumerate(row_values):
        if i > 0:
            draw.line([(x, y), (x, y + row_h)], fill=(0, 0, 0), width=2)
        tw = draw.textlength(val, font=font_body)
        draw.text((x + (col_widths[i] - tw) / 2, y + 10), val, fill=(0, 0, 0), font=font_body)
        x += col_widths[i]

    # 总金额行
    y += row_h
    draw.rectangle([(x0, y), (x0 + table_w, y + row_h)], outline=(0, 0, 0), width=2)
    draw.line([(x0 + col_widths[0], y), (x0 + col_widths[0], y + row_h)], fill=(0, 0, 0), width=2)
    draw.text((x0 + 34, y + 10), "总金额", fill=(0, 0, 0), font=font_head)
    total_text = f"小写：{amount:.2f}，大写：{amount_upper}"
    draw.text((x0 + col_widths[0] + 18, y + 10), total_text, fill=(0, 0, 0), font=font_body)

    # 配置说明行（动态高度）
    config_label_w = col_widths[0]
    config_content_w = table_w - config_label_w - 26
    wrapped_lines: List[str] = []
    for line in config_lines:
        wrapped_lines.extend(wrap_text_by_width(draw, line, font_body, config_content_w))
    config_h = max(170, len(wrapped_lines) * 32 + 24)

    y += row_h
    draw.rectangle([(x0, y), (x0 + table_w, y + config_h)], outline=(0, 0, 0), width=2)
    draw.line([(x0 + config_label_w, y), (x0 + config_label_w, y + config_h)], fill=(0, 0, 0), width=2)
    draw.text((x0 + 35, y + config_h / 2 - 12), "配置", fill=(0, 0, 0), font=font_head)

    text_y = y + 12
    for line in wrapped_lines:
        draw.text((x0 + config_label_w + 14, text_y), line, fill=(0, 0, 0), font=font_body)
        text_y += 32

    # 条款
    y += config_h + 24
    note_lines = [
        "2.以上价格：含税，不含运输费用。",
        f"3.付款方式：{form_data.get('payment_term', '')}",
        f"4.质保期：{form_data.get('warranty_term', '')}",
        f"5.交货期：{form_data.get('delivery_term', '')}",
        f"6.有效期：{form_data.get('validity_term', '')}",
    ]
    for note in note_lines:
        lines = wrap_text_by_width(draw, note, font_note, 1020)
        for line in lines:
            draw.text((100, y), line, fill=(0, 0, 0), font=font_note)
            y += 34

    # 左下角公司信息（可在后台勾选显示/隐藏）
    footer_lines: List[str] = []
    show_company = str(settings.get("quotation_show_company", "1")).strip() == "1"
    show_contact = str(settings.get("quotation_show_contact", "1")).strip() == "1"
    company_name = str(settings.get("company_name", "")).strip()
    company_contact = str(settings.get("company_contact", "")).strip()
    company_phone = str(settings.get("company_phone", "")).strip()

    if show_company and company_name:
        footer_lines.append(f"公司名称：{company_name}")
    if show_contact:
        contact_parts = [part for part in [company_contact, company_phone] if part]
        if contact_parts:
            footer_lines.append(f"联系方式：{' / '.join(contact_parts)}")

    if footer_lines:
        start_y = 1540 if len(footer_lines) > 1 else 1572
        for idx, line in enumerate(footer_lines):
            draw.text((100, start_y + idx * 34), line, fill=(0, 0, 0), font=font_note)

    # 日期
    draw.text((860, 1600), f"报价日期：{datetime.now().strftime('%Y.%m.%d')}", fill=(0, 0, 0), font=font_note)

    image.save(out_file, "PNG")
    return rel


def generate_excel(quote: Dict[str, Any]) -> Path:
    out_file = GENERATED_DIR / f"quotation_{quote['quote_no']}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    ws["A1"] = f"报价单 #{quote['quote_no']}"
    ws["A2"] = f"生成时间：{quote['created_at']}"
    ws.append([])
    ws.append(["序号", "配置项", "价格"])
    for index, item in enumerate(quote["selections"], start=1):
        ws.append([index, f"{item['category_name']}：{item['option_name']}", float(item["price"])])
    ws.append([])
    ws.append(["", "未税小计", quote["subtotal"]])
    ws.append(["", f"税额({quote['tax_rate']*100:.0f}%)", quote["tax_amount"]])
    ws.append(["", "含税总价", quote["total"]])
    wb.save(out_file)
    return out_file


def generate_params_image(items: List[Dict[str, Any]]) -> str:
    line_height = 34
    head_height = 90
    footer_height = 40
    content_lines = 0
    for item in items:
        params = item.get("params", [])
        content_lines += 1 + max(1, len(params))
    height = max(520, head_height + footer_height + content_lines * line_height + 80)

    out_file = GENERATED_DIR / f"params_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.png"
    rel = f"/generated/{out_file.name}"

    image = Image.new("RGB", (1240, height), "white")
    draw = ImageDraw.Draw(image)
    font_title = load_font(44, True)
    font_sub = load_font(24, True)
    font_body = load_font(20)

    draw.rectangle([(30, 30), (1210, 110)], fill=(241, 247, 255), outline=(60, 90, 140), width=2)
    title = "所选产品参数清单"
    title_w = draw.textlength(title, font=font_title)
    draw.text(((1240 - title_w) / 2, 42), title, fill=(28, 63, 112), font=font_title)

    y = 150
    for index, item in enumerate(items, start=1):
        draw.text(
            (60, y),
            f"{index}. {item.get('category_name', '')}：{item.get('option_name', '')}（¥{item.get('price', 0):,.2f}）",
            fill=(32, 76, 145),
            font=font_sub,
        )
        y += line_height

        params = item.get("params", [])
        if not params:
            draw.text((90, y), "参数：未填写", fill=(60, 60, 60), font=font_body)
            y += line_height
        else:
            for param in params:
                draw.text(
                    (90, y),
                    f"- {param.get('key', '参数')}：{param.get('value', '')}",
                    fill=(60, 60, 60),
                    font=font_body,
                )
                y += line_height
        y += 10

    image.save(out_file, "PNG")
    return rel


@app.route("/")
@app.route("/index")
@app.route("/index.html")
@app.route("/new_app/pages/price_calculator_case/index")
def index():
    share_meta = build_share_meta()
    return render_template(
        "index.html",
        title=share_meta["title"],
        share_meta=share_meta,
        wechat_app_id=WECHAT_APP_ID,
    )
    return render_template(
        "index.html",
        title="发电机组报价系统",
        share_meta=share_meta,
        wechat_app_id=WECHAT_APP_ID,
    )
    return render_template("index.html", title="发电机组报价系统")


@app.get("/params")
def params_page():
    return render_template("params.html", title="产品参数")


@app.get("/health")
def health():
    return jsonify({"ok": True, "time": datetime.now().isoformat()})


@app.get("/generated/<path:filename>")
def generated_file(filename: str):
    safe_name = Path(filename).name
    abs_path = GENERATED_DIR / safe_name
    if not abs_path.exists():
        return "Not Found", 404
    return send_file(abs_path, as_attachment=False)


@app.get("/api/config")
def api_config():
    return jsonify({"categories": categories_with_options(True), "settings": get_settings()})


@app.get("/api/wechat/jssdk-config")
def api_wechat_jssdk_config():
    if not WECHAT_APP_ID or not WECHAT_APP_SECRET:
        return jsonify({"ok": False, "message": "未配置微信公众号 AppID/AppSecret"}), 503

    target_url = str(request.args.get("url", "")).strip()
    if not target_url:
        return jsonify({"ok": False, "message": "缺少 url 参数"}), 400

    try:
        config = build_wechat_signature(target_url)
    except Exception as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400

    return jsonify({"ok": True, "config": config, "share": build_share_meta()})


@app.post("/api/quotes")
def api_quotes():
    payload = request.get_json(silent=True) or {}
    selection_map = parse_selections(payload.get("selections", {}))
    items, errors = build_quote_items(selection_map)
    if errors:
        return jsonify({"ok": False, "message": "；".join(errors)}), 400
    if not items:
        return jsonify({"ok": False, "message": "请至少选择一个配置项"}), 400

    include_tax = bool(payload.get("include_tax", True))
    base_tax_rate = max(0.0, to_float(get_settings().get("tax_rate", "0.13"), 0.13))
    tax_rate = base_tax_rate if include_tax else 0.0
    subtotal = round(sum(item["price"] for item in items), 2)
    tax_amount = round(subtotal * tax_rate, 2)
    total = round(subtotal + tax_amount, 2)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    db = get_db()
    cur = db.execute(
        """
        INSERT INTO quotes(
            quote_no,customer_name,customer_phone,company_name,project_name,notes,
            subtotal,tax_rate,tax_amount,total,selections_json,created_at
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            generate_quote_no(),
            str(payload.get("customer_name", "")).strip(),
            str(payload.get("customer_phone", "")).strip(),
            str(payload.get("company_name", "")).strip(),
            str(payload.get("project_name", "")).strip(),
            str(payload.get("notes", "")).strip(),
            subtotal,
            tax_rate,
            tax_amount,
            total,
            json.dumps(items, ensure_ascii=False),
            now,
        ),
    )
    db.commit()

    quote_id = cur.lastrowid
    quote_no = db.execute("SELECT quote_no FROM quotes WHERE id=?", (quote_id,)).fetchone()["quote_no"]
    quote_data = row_quote(quote_id) or {
        "selections": items,
        "total": total,
        "quote_no": quote_no,
        "created_at": now,
    }
    quotation_form = build_quotation_form_defaults(quote_data)
    return jsonify(
        {
            "ok": True,
            "quote_id": quote_id,
            "quote_no": quote_no,
            "subtotal": subtotal,
            "tax_rate": tax_rate,
            "tax_amount": tax_amount,
            "total": total,
            "include_tax": include_tax,
            "created_at": now,
            "contract_image_url": f"/api/quotes/{quote_id}/contract-image",
            "quotation_image_url": f"/api/quotes/{quote_id}/quotation-image",
            "quotation_excel_url": f"/api/quotes/{quote_id}/quotation-excel",
            "quotation_url": f"/api/quotes/{quote_id}/quotation",
            "quotation_form": quotation_form,
        }
    )


@app.post("/api/parameter-image")
def api_parameter_image():
    payload = request.get_json(silent=True) or {}
    selection_map = parse_selections(payload.get("selections", {}))
    items, errors = build_quote_items(selection_map, strict_required=False)
    if errors:
        return jsonify({"ok": False, "message": "；".join(errors)}), 400
    if not items:
        return jsonify({"ok": False, "message": "请先选择商品"}), 400

    rel = generate_params_image(items)
    return jsonify({"ok": True, "image_url": f"{rel}?t={int(datetime.now().timestamp())}"})


@app.get("/api/quotes/<int:quote_id>/contract-image")
def api_contract_image(quote_id: int):
    quote = row_quote(quote_id)
    if not quote:
        return jsonify({"ok": False, "message": "报价不存在"}), 404
    rel = (quote.get("contract_image_path") or "").strip()
    abs_path = resolve_generated_path(rel)
    if not rel or not abs_path or not abs_path.exists():
        rel = generate_contract_image(quote)
        db = get_db()
        db.execute("UPDATE quotes SET contract_image_path=? WHERE id=?", (rel, quote_id))
        db.commit()
    return jsonify(
        {
            "ok": True,
            "quote_id": quote_id,
            "quote_no": quote["quote_no"],
            "image_url": f"{rel}?t={int(datetime.now().timestamp())}",
        }
    )


@app.get("/api/quotes/<int:quote_id>/quotation")
def api_quotation(quote_id: int):
    quote = row_quote(quote_id)
    if not quote:
        return jsonify({"ok": False, "message": "报价不存在"}), 404
    form_data = merge_quotation_form(build_quotation_form_defaults(quote), quote.get("quotation_meta", {}))
    rel = (quote.get("quotation_image_path") or "").strip()
    abs_path = resolve_generated_path(rel)
    if not rel or not abs_path or not abs_path.exists():
        rel = generate_quotation_image(quote, form_data)
        abs_path = resolve_generated_path(rel)
        db = get_db()
        db.execute("UPDATE quotes SET quotation_image_path=? WHERE id=?", (rel, quote_id))
        db.commit()
    return send_file(abs_path, mimetype="image/png", as_attachment=False, download_name=f"{quote['quote_no']}_报价单.png")


@app.route("/api/quotes/<int:quote_id>/quotation-image", methods=["GET", "POST"])
def api_quotation_image(quote_id: int):
    quote = row_quote(quote_id)
    if not quote:
        return jsonify({"ok": False, "message": "报价不存在"}), 404

    base_form = merge_quotation_form(build_quotation_form_defaults(quote), quote.get("quotation_meta", {}))

    if request.method == "POST":
        payload = request.get_json(silent=True) or {}
        form_data = merge_quotation_form(base_form, payload.get("form", {}))
        rel = generate_quotation_image(quote, form_data)
        db = get_db()
        db.execute(
            "UPDATE quotes SET quotation_image_path=?, quotation_meta_json=? WHERE id=?",
            (rel, json.dumps(form_data, ensure_ascii=False), quote_id),
        )
        db.commit()
        return jsonify(
            {
                "ok": True,
                "quote_id": quote_id,
                "quote_no": quote["quote_no"],
                "image_url": f"{rel}?t={int(datetime.now().timestamp())}",
                "form": form_data,
            }
        )

    rel = (quote.get("quotation_image_path") or "").strip()
    abs_path = resolve_generated_path(rel)
    if not rel or not abs_path or not abs_path.exists():
        rel = generate_quotation_image(quote, base_form)
        db = get_db()
        db.execute(
            "UPDATE quotes SET quotation_image_path=?, quotation_meta_json=? WHERE id=?",
            (rel, json.dumps(base_form, ensure_ascii=False), quote_id),
        )
        db.commit()
    return jsonify(
        {
            "ok": True,
            "quote_id": quote_id,
            "quote_no": quote["quote_no"],
            "image_url": f"{rel}?t={int(datetime.now().timestamp())}",
            "form": base_form,
        }
    )


@app.get("/api/quotes/<int:quote_id>/quotation-excel")
def api_quotation_excel(quote_id: int):
    quote = row_quote(quote_id)
    if not quote:
        return jsonify({"ok": False, "message": "报价不存在"}), 404
    rel = (quote.get("quotation_file_path") or "").strip()
    abs_path = resolve_generated_path(rel)
    if not rel or not abs_path or not abs_path.exists():
        abs_path = generate_excel(quote)
        rel = f"/generated/{abs_path.name}"
        db = get_db()
        db.execute("UPDATE quotes SET quotation_file_path=? WHERE id=?", (rel, quote_id))
        db.commit()
    return send_file(abs_path, as_attachment=True, download_name=f"{quote['quote_no']}_报价单.xlsx")


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if ADMIN_PASS and username == ADMIN_USER and password == ADMIN_PASS:
            session["is_admin"] = True
            return redirect(url_for("admin_dashboard"))
        flash("账号或密码错误")
    return render_template("admin_login.html", title="后台登录")


@app.get("/admin/logout")
@admin_required
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))


@app.get("/admin")
@admin_required
def admin_dashboard():
    return render_template(
        "admin_dashboard.html",
        title="后台管理",
        categories=categories_with_options(False),
        settings=get_settings(),
    )


@app.post("/admin/settings")
@admin_required
def admin_update_settings():
    set_settings(
        {
            "company_name": request.form.get("company_name", "").strip(),
            "company_contact": request.form.get("company_contact", "").strip(),
            "company_phone": request.form.get("company_phone", "").strip(),
            "quotation_show_company": "1" if request.form.get("quotation_show_company") == "1" else "0",
            "quotation_show_contact": "1" if request.form.get("quotation_show_contact") == "1" else "0",
            "company_address": request.form.get("company_address", "").strip(),
            "tax_rate": request.form.get("tax_rate", "0.13").strip(),
            "contract_notice": request.form.get("contract_notice", "").strip(),
            "delivery_term": request.form.get("delivery_term", "").strip(),
        }
    )
    flash("系统设置已更新")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/categories/add")
@admin_required
def admin_add_category():
    name = request.form.get("name", "").strip()
    if not name:
        flash("分类名称不能为空")
        return redirect(url_for("admin_dashboard"))
    db = get_db()
    try:
        db.execute(
            """
            INSERT INTO categories(name,description,required,selection_mode,sort_order)
            VALUES(?,?,?,?,?)
            """,
            (
                name,
                request.form.get("description", "").strip(),
                1 if request.form.get("required") == "1" else 0,
                normalize_mode(request.form.get("selection_mode")),
                to_int(request.form.get("sort_order"), 0),
            ),
        )
        db.commit()
        flash("分类已新增")
    except sqlite3.IntegrityError:
        flash("分类名称重复")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/categories/<int:category_id>/update")
@admin_required
def admin_update_category(category_id: int):
    name = request.form.get("name", "").strip()
    if not name:
        flash("分类名称不能为空")
        return redirect(url_for("admin_dashboard"))
    db = get_db()
    try:
        db.execute(
            """
            UPDATE categories
            SET name=?,description=?,required=?,selection_mode=?,sort_order=?
            WHERE id=?
            """,
            (
                name,
                request.form.get("description", "").strip(),
                1 if request.form.get("required") == "1" else 0,
                normalize_mode(request.form.get("selection_mode")),
                to_int(request.form.get("sort_order"), 0),
                category_id,
            ),
        )
        db.commit()
        flash("分类已更新")
    except sqlite3.IntegrityError:
        flash("分类名称重复")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/categories/<int:category_id>/delete")
@admin_required
def admin_delete_category(category_id: int):
    db = get_db()
    db.execute("DELETE FROM categories WHERE id=?", (category_id,))
    db.commit()
    flash("分类已删除")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/options/add")
@admin_required
def admin_add_option():
    name = request.form.get("name", "").strip()
    if not name:
        flash("选项名称不能为空")
        return redirect(url_for("admin_dashboard"))
    category_id = to_int(request.form.get("category_id"), 0)
    if category_id <= 0:
        flash("请选择所属分类")
        return redirect(url_for("admin_dashboard"))
    db = get_db()
    try:
        db.execute(
            """
            INSERT INTO options(category_id,name,description,price,image_path,params_text,sort_order,is_active)
            VALUES(?,?,?,?,?,?,?,?)
            """,
            (
                category_id,
                name,
                request.form.get("description", "").strip(),
                to_float(request.form.get("price"), 0),
                request.form.get("image_path", "").strip(),
                request.form.get("params_text", "").strip(),
                to_int(request.form.get("sort_order"), 0),
                1 if request.form.get("is_active", "1") == "1" else 0,
            ),
        )
        db.commit()
        flash("选项已新增")
    except sqlite3.IntegrityError:
        flash("新增失败，请检查分类是否存在")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/options/<int:option_id>/update")
@admin_required
def admin_update_option(option_id: int):
    name = request.form.get("name", "").strip()
    if not name:
        flash("选项名称不能为空")
        return redirect(url_for("admin_dashboard"))
    db = get_db()
    db.execute(
        """
        UPDATE options
        SET name=?,description=?,price=?,image_path=?,params_text=?,sort_order=?,is_active=?
        WHERE id=?
        """,
        (
            name,
            request.form.get("description", "").strip(),
            to_float(request.form.get("price"), 0),
            request.form.get("image_path", "").strip(),
            request.form.get("params_text", "").strip(),
            to_int(request.form.get("sort_order"), 0),
            1 if request.form.get("is_active", "1") == "1" else 0,
            option_id,
        ),
    )
    db.commit()
    flash("选项已更新")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/options/<int:option_id>/delete")
@admin_required
def admin_delete_option(option_id: int):
    db = get_db()
    db.execute("DELETE FROM options WHERE id=?", (option_id,))
    db.commit()
    flash("选项已删除")
    return redirect(url_for("admin_dashboard"))


@app.post("/admin/options/<int:option_id>/upload-image")
@admin_required
def admin_upload_option_image(option_id: int):
    image_file = request.files.get("image")
    if not image_file or not image_file.filename:
        flash("请选择图片")
        return redirect(url_for("admin_dashboard"))

    ext = Path(image_file.filename).suffix.lower()
    if ext not in ALLOWED_IMAGE_EXT:
        flash("只支持 png/jpg/jpeg/webp")
        return redirect(url_for("admin_dashboard"))

    filename = f"{uuid.uuid4().hex}{ext}"
    image_file.save(UPLOAD_DIR / filename)

    db = get_db()
    db.execute("UPDATE options SET image_path=? WHERE id=?", (f"/static/uploads/{filename}", option_id))
    db.commit()
    flash("图片上传成功")
    return redirect(url_for("admin_dashboard"))


@app.errorhandler(404)
def handle_404(_e):
    path = request.path or ""
    if path.startswith("/api/"):
        return jsonify({"ok": False, "message": "API not found"}), 404
    if path.startswith("/admin"):
        return redirect(url_for("admin_login"))
    if path.startswith("/static/"):
        return "Not Found", 404
    return redirect(url_for("index"))


with app.app_context():
    init_db()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
